"""Template loaders.

Two file types are supported:

* ``.docx``  - parsed with ``python-docx``. Lossless paragraph extraction.
* ``.doc``   - legacy binary Word. We use pure-Python OLE parsing
              (``olefile``) to pull text from the ``WordDocument`` stream.
              This is rough but requires no external binary (no LibreOffice,
              no MS Word). Quality is good enough for RAG context: we don't
              need byte-perfect fidelity, we only need readable Arabic text
              the embedding model can match on.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import olefile
from docx import Document


@dataclass
class LoadedTemplate:
    """Parsed template ready to be chunked / indexed."""

    source: str          # file name only, used for citation
    category: str        # CONTRACT_CATEGORIES key
    text: str            # cleaned full text


# ---------------------------------------------------------------------------
# .docx
# ---------------------------------------------------------------------------
def _read_docx(path: Path) -> str:
    """Extract paragraph + table text from a .docx file."""
    doc = Document(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# .doc (legacy OLE compound document)
# ---------------------------------------------------------------------------
# The .doc format is an OLE2 compound file. The actual text lives in the
# ``WordDocument`` stream interleaved with formatting control bytes. The
# fastest "good enough" recovery is to:
#   1. read the stream as bytes,
#   2. try decoding as UTF-16-LE (most modern Arabic .doc files), and as
#      CP-1256 (older Arabic Windows encoding) as a fallback,
#   3. strip control characters,
#   4. keep runs of printable text long enough to be meaningful.
_PRINTABLE_RE = re.compile(
    r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF"  # Arabic
    r"a-zA-Z0-9"
    r"\s"
    r"\.,;:!?\-–—()«»\"'/\\&%@#\u060C\u061B\u061F]+"  # common punctuation incl. Arabic comma/semi/?
)

_CONTROL_RE = re.compile(r"[\x00-\x08\x0B-\x1F\x7F]")

# Word puts these word-art / field markers inline; we drop them.
_FIELD_NOISE_RE = re.compile(
    r"(HYPERLINK|PAGE\s*\\\*|TOC|MERGEFORMAT|FORMTEXT|FORMCHECKBOX|MACROBUTTON)"
)


def _decode_word_stream(data: bytes) -> str:
    """Best-effort text recovery from a raw WordDocument byte stream."""
    candidates: list[str] = []

    # Try UTF-16-LE first (Word 97+ stores Unicode for most Arabic content)
    try:
        utf16 = data.decode("utf-16-le", errors="ignore")
        candidates.append(utf16)
    except Exception:
        pass

    # Try CP-1256 (older Arabic Windows files sometimes still appear)
    try:
        cp1256 = data.decode("cp1256", errors="ignore")
        candidates.append(cp1256)
    except Exception:
        pass

    # Pick whichever yielded more Arabic letters
    def arabic_score(s: str) -> int:
        return sum(1 for ch in s if "\u0600" <= ch <= "\u06FF")

    if not candidates:
        return ""
    return max(candidates, key=arabic_score)


def _clean_extracted(text: str) -> str:
    """Drop control chars / field noise, normalize whitespace."""
    text = _CONTROL_RE.sub(" ", text)
    text = _FIELD_NOISE_RE.sub(" ", text)
    text = unicodedata.normalize("NFKC", text)

    # Keep only runs of meaningful printable characters
    kept: list[str] = []
    for match in _PRINTABLE_RE.finditer(text):
        chunk = match.group(0).strip()
        if len(chunk) >= 8:  # discard tiny fragments
            kept.append(chunk)

    cleaned = "\n".join(kept)
    # Collapse runs of whitespace
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _read_doc(path: Path) -> str:
    """Extract readable text from a legacy .doc file (pure Python)."""
    try:
        ole = olefile.OleFileIO(str(path))
    except Exception:
        return ""

    try:
        if not ole.exists("WordDocument"):
            return ""
        raw = ole.openstream("WordDocument").read()
    finally:
        ole.close()

    decoded = _decode_word_stream(raw)
    return _clean_extracted(decoded)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def load_template(path: Path, category: str) -> LoadedTemplate | None:
    """Load one template file. Returns None on failure or empty content."""
    if not path.exists():
        return None

    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            text = _read_docx(path)
        elif suffix == ".doc":
            text = _read_doc(path)
        else:
            return None
    except Exception:
        return None

    text = text.strip()
    if len(text) < 100:
        # Too short to be a useful contract template
        return None

    return LoadedTemplate(source=path.name, category=category, text=text)


def load_all_templates(
    templates_root: Path,
    categories: dict[str, list[str]],
) -> list[LoadedTemplate]:
    """Load every template referenced in ``categories``.

    Missing files are silently skipped (so the app still runs if a few are
    moved/renamed).
    """
    loaded: list[LoadedTemplate] = []
    seen: set[str] = set()

    for category, files in categories.items():
        for fname in files:
            if fname in seen:
                continue
            seen.add(fname)
            tpl = load_template(templates_root / fname, category)
            if tpl is not None:
                loaded.append(tpl)

    return loaded
